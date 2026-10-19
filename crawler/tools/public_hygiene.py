"""Public-fixture hygiene checks shared by tests and history scanning."""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pypdf import PdfReader


MAX_FIXTURE_BYTES = 512 * 1024
MAX_FIXTURE_TOTAL_BYTES = 1024 * 1024
MAX_XLSX_ROWS = 30
MAX_XLSX_COLUMNS = 16
MAX_PDF_PAGES = 4
MAX_HTML_BYTES = 64 * 1024
SYNTHETIC_MARKER = "synthetic-fixture"
SYNTHETIC_AUTHOR = "Synthetic Fixture Generator"

PERSONAL_PATTERNS = {
    "personal-phone": re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)"),
    "personal-email": re.compile(
        r"(?<![A-Za-z0-9._%+-])[A-Za-z0-9._%+-]+@"
        r"[A-Za-z0-9.-]+\.[A-Za-z]{2,}"
    ),
    "personal-id": re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)"),
    "landline-phone": re.compile(r"(?<!\d)0\d{2,3}[- ]?\d{7,8}(?!\d)"),
}
TOKEN_PATTERN = re.compile(
    r"(?i)(__VIEWSTATE|__EVENTVALIDATION|__RequestVerificationToken|"
    r"csrf|authenticity_token|access_token|api[_-]?key|bearer\s+\S+)"
)
CONTACT_LABELS = (
    "联系人",
    "联系电话",
    "手机号",
    "电话",
    "邮箱",
    "电子邮箱",
    "身份证",
    "微信号",
    "QQ",
)
CONTACT_FIELD_PATTERN = re.compile(
    r"(?:"
    + "|".join(re.escape(label) for label in CONTACT_LABELS)
    + r")\s*[:：=]\s*([^\s|,，;；<>]{1,40})",
    re.IGNORECASE,
)
SENSITIVE_HEADER_PATTERN = re.compile(
    "|".join(re.escape(label) for label in CONTACT_LABELS),
    re.IGNORECASE,
)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8", "gb18030", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _text_findings(text: str) -> set[str]:
    findings: set[str] = set()
    for category, pattern in PERSONAL_PATTERNS.items():
        if pattern.search(text):
            findings.add(category)
    if TOKEN_PATTERN.search(text):
        findings.add("hidden-form-token")
    if CONTACT_FIELD_PATTERN.search(text):
        findings.add("personal-contact-field")
    if re.search(r"(?i)\b(?:mailto|tel):", text):
        findings.add("personal-contact-field")
    return findings


def _size_findings(data: bytes) -> set[str]:
    return {"oversized-fixture"} if len(data) > MAX_FIXTURE_BYTES else set()


def _scan_html(data: bytes) -> set[str]:
    findings = _size_findings(data)
    if len(data) > MAX_HTML_BYTES:
        findings.add("full-source-attachment")
    soup = BeautifulSoup(data, "html.parser")
    marker = soup.find(
        "meta",
        attrs={"name": re.compile(r"^fixture-origin$", re.IGNORECASE)},
    )
    marker_value = (marker.get("content") if marker else "") or ""
    if marker_value.lower() != SYNTHETIC_MARKER:
        findings.add("missing-synthetic-marker")

    hidden_inputs = soup.find_all(
        "input", attrs={"type": re.compile(r"^hidden$", re.IGNORECASE)}
    )
    if soup.find("form"):
        findings.add("original-page-form")
    if soup.find("script") or soup.find("iframe"):
        findings.add("original-page-embed")
    if hidden_inputs:
        findings.add("hidden-form-state")
    for node in hidden_inputs:
        name = str(node.get("name") or "")
        if TOKEN_PATTERN.search(name):
            findings.add("hidden-form-token")

    findings.update(_text_findings(soup.get_text(" ", strip=True)))
    for node in soup.find_all(href=True):
        href = str(node.get("href") or "")
        if re.match(r"(?i)^(?:mailto|tel):", href):
            findings.add("personal-contact-field")
        findings.update(_text_findings(href))
    return findings


def _xlsx_metadata(zf: zipfile.ZipFile) -> tuple[str, str, str]:
    core = ""
    app = ""
    try:
        core = zf.read("docProps/core.xml").decode("utf-8", errors="ignore")
    except KeyError:
        pass
    try:
        app = zf.read("docProps/app.xml").decode("utf-8", errors="ignore")
    except KeyError:
        pass

    def field(xml_text: str, local_name: str) -> str:
        try:
            root = ElementTree.fromstring(xml_text)
        except ElementTree.ParseError:
            return ""
        for node in root.iter():
            if node.tag.rsplit("}", 1)[-1] == local_name:
                return (node.text or "").strip()
        return ""

    return (
        field(core, "creator"),
        field(core, "lastModifiedBy"),
        field(app, "Company"),
    )


def _scan_xlsx(data: bytes) -> set[str]:
    findings = _size_findings(data)
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        return findings | {"invalid-fixture"}

    with zf:
        names = set(zf.namelist())
        if any(
            name.startswith(("xl/media/", "xl/embeddings/", "xl/externalLinks/"))
            or name.endswith("vbaProject.bin")
            for name in names
        ):
            findings.add("embedded-attachment")
        try:
            core = zf.read("docProps/core.xml").decode(
                "utf-8", errors="ignore"
            )
        except KeyError:
            core = ""
        if SYNTHETIC_MARKER not in core.lower():
            findings.add("missing-synthetic-marker")
        creator, last_modified_by, company = _xlsx_metadata(zf)
        for value in (creator, last_modified_by, company):
            if value and value.lower() != SYNTHETIC_AUTHOR.lower():
                findings.add("original-attachment-metadata")
                break
        for name in names:
            if name.endswith(".xml") or name.endswith(".rels"):
                findings.update(_text_findings(_decode_text(zf.read(name))))

    try:
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=True
        )
    except Exception:
        return findings | {"invalid-fixture"}
    try:
        for sheet in workbook.worksheets:
            if sheet.max_row > MAX_XLSX_ROWS or sheet.max_column > MAX_XLSX_COLUMNS:
                findings.add("full-source-attachment")
            rows = list(sheet.iter_rows(values_only=True))
            for row in rows[:10]:
                headers = [str(value or "").strip() for value in row]
                contact_columns = [
                    index
                    for index, value in enumerate(headers)
                    if SENSITIVE_HEADER_PATTERN.search(value)
                ]
                if not contact_columns:
                    continue
                for data_row in rows[1:]:
                    if any(
                        index < len(data_row)
                        and str(data_row[index] or "").strip()
                        for index in contact_columns
                    ):
                        findings.add("personal-contact-column")
                        break
                break
            for row in rows:
                findings.update(
                    _text_findings(
                        " ".join(str(value) for value in row if value is not None)
                    )
                )
    finally:
        workbook.close()
    return findings


def _scan_pdf(data: bytes) -> set[str]:
    findings = _size_findings(data)
    if b"/EmbeddedFiles" in data or b"/Filespec" in data:
        findings.add("embedded-attachment")
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
    except Exception:
        return findings | {"invalid-fixture"}
    metadata = reader.metadata or {}
    marker_text = " ".join(
        str(metadata.get(key) or "")
        for key in ("/Title", "/Subject", "/Keywords")
    )
    if SYNTHETIC_MARKER not in marker_text.lower():
        findings.add("missing-synthetic-marker")
    author = str(metadata.get("/Author") or "").strip()
    creator = str(metadata.get("/Creator") or "").strip()
    producer = str(metadata.get("/Producer") or "").strip()
    for value in (author, creator, producer):
        if value and SYNTHETIC_AUTHOR.lower() not in value.lower():
            findings.add("original-attachment-metadata")
            break
    if len(reader.pages) > MAX_PDF_PAGES:
        findings.add("full-source-attachment")
    extracted = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        try:
            extracted.append(page.extract_text() or "")
        except Exception:
            extracted.append("")
    text = "\n".join(extracted)
    if SYNTHETIC_MARKER not in text.lower():
        findings.add("missing-synthetic-marker")
    if re.search(r"[\u4e00-\u9fff]", text):
        findings.add("original-source-text")
    findings.update(_text_findings(text))
    return findings


def scan_fixture_bytes(path: str | Path, data: bytes) -> list[str]:
    """Return hygiene categories for one public evidence fixture."""
    suffix = Path(path).suffix.lower()
    if suffix == ".html" or suffix == ".htm":
        findings = _scan_html(data)
    elif suffix == ".xlsx":
        findings = _scan_xlsx(data)
    elif suffix == ".pdf":
        findings = _scan_pdf(data)
    else:
        findings = _size_findings(data)
        findings.update(_text_findings(_decode_text(data)))
    return sorted(findings)
