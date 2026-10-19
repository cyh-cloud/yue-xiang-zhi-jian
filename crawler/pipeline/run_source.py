import json
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from urllib.error import HTTPError
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from . import database, source_registry
from .evidence import (
    RawEvidence,
    guess_content_type,
    persist_raw_evidence,
    sha256_bytes,
    validate_detail_locator,
)
from .publication import evaluate_publication
from .sources import cass_linguistics


DEFAULT_MAX_DOWNLOAD_BYTES = 100 * 1024 * 1024

_TEST_SOURCE_MODES = {
    "fixture-source": frozenset({"fixture"}),
}
_SOURCE_MODES = {
    source_id: source_registry.run_modes(source)
    for source_id, source in source_registry.load_registry().items()
}
_SOURCE_MODES.update(_TEST_SOURCE_MODES)
RUN_MODES = tuple(
    sorted({mode for modes in _SOURCE_MODES.values() for mode in modes})
)

_AQSC_FILES = {
    "synthetic-mtyx-sample.xlsx": {
        "cert_type": "名特优新",
        "announcement_url": (
            "http://www.aqsc.agri.cn/tzgg/202511/t20251128_468194.htm"
        ),
        "attachment_url": (
            "http://www.aqsc.agri.cn/tzgg/202511/"
            "P020251225613822033339.xlsx"
        ),
    },
    "synthetic-tepinongpin-sample.xlsx": {
        "cert_type": "特质农品",
        "announcement_url": (
            "http://www.aqsc.agri.cn/tzgg/202511/t20251128_468195.htm"
        ),
        "attachment_url": (
            "http://www.aqsc.agri.cn/tzgg/202511/"
            "P020251128359343750170.xlsx"
        ),
    },
}

_CNIPA_DETAIL_URLS = {
    "687": "https://ggfw.cnipa.gov.cn/dlbzsq/detail?id=541",
    "89": "https://ggfw.cnipa.gov.cn/dlbzsq/detail?id=58",
    "539": "https://ggfw.cnipa.gov.cn/dlbzsq/detail?id=513",
}


@dataclass(frozen=True)
class RunConfig:
    database_path: Path
    download_root: Path
    fixture_root: Path
    page_limit: int = 100
    xlsx_root: Path | None = None
    pdf_root: Path | None = None
    source_root: Path | None = None
    registry_path: Path | None = None
    max_download_bytes: int = DEFAULT_MAX_DOWNLOAD_BYTES


@dataclass(frozen=True)
class SourceRunResult:
    run_id: str
    source_id: str
    status: str
    discovered_count: int
    fetched_count: int
    written_count: int
    skipped_count: int
    failed_count: int
    errors: tuple[str, ...] = ()


@dataclass
class DiscoveryResult:
    items: list = field(default_factory=list)
    evidence: list[RawEvidence] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def _now():
    return datetime.now(timezone.utc).isoformat()


def _json_bytes(value):
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")


def _fixture_path(config):
    return Path(config.fixture_root) / "discover.json"


def _fixture_metadata(path):
    body = path.read_bytes()
    return body, {
        "fixture_path": str(path),
        "fixture_sha256": sha256_bytes(body),
    }


def _evidence(
    source_id,
    key,
    *,
    request_url,
    source_url,
    long_term_url,
    body,
    content_type,
    detail_url=None,
    detail_url_status="missing",
    fetch_url=None,
    locator=None,
    file_name=None,
    rights_status="unknown",
    attribution_required=False,
    attribution_text=None,
    rights_evidence_url=None,
    status_code=200,
    response_headers=None,
    metadata=None,
    fetch_status="ok",
    error_message=None,
    declared_size=None,
    final_url=None,
    file_url=None,
):
    return RawEvidence(
        source_id=source_id,
        key=key,
        request_url=request_url,
        final_url=final_url or request_url,
        source_url=source_url,
        long_term_url=long_term_url,
        detail_url=detail_url,
        detail_url_status=detail_url_status,
        fetch_url=fetch_url,
        locator=locator,
        file_url=file_url,
        body=body,
        content_type=content_type,
        file_name=file_name,
        status_code=status_code,
        request_headers={
            "user-agent": "yuexiang-crawler/0.1 (competition; noncommercial)"
        },
        response_headers=response_headers or {"content-type": content_type},
        rights_status=rights_status,
        attribution_required=attribution_required,
        attribution_text=attribution_text,
        rights_evidence_url=rights_evidence_url,
        fetch_status=fetch_status,
        error_message=error_message,
        declared_size=declared_size,
        metadata={"source_id": source_id, **(metadata or {})},
    )


def _rights_fields(source_id, config=None):
    registry = source_registry.load_registry(
        getattr(config, "registry_path", None)
    )
    source = registry.get(source_id)
    if source is not None:
        return source_registry.rights_fields(source)
    if source_id == "fixture-source":
        return {
            "rights_status": "noncommercial-permitted",
            "attribution_required": True,
            "attribution_text": "Example source",
            "rights_evidence_url": "https://example.test/copyright",
        }
    return {
        "rights_status": "unknown",
        "attribution_required": False,
        "attribution_text": None,
        "rights_evidence_url": None,
    }


def _discover_generic_fixture(source_id, config):
    path = _fixture_path(config)
    if not path.exists():
        return DiscoveryResult(errors=[f"fixture not found: {path}"])
    body, fixture_meta = _fixture_metadata(path)
    payload = json.loads(body.decode("utf-8"))
    raw_items = payload.get("items") or []
    evidence = []
    items = []
    if not raw_items:
        return DiscoveryResult(errors=[f"fixture has no items: {path}"])
    list_url = raw_items[0].get("source_url") or "https://fixture.local/list"
    list_long_term = raw_items[0].get("long_term_url") or list_url
    rights = _rights_fields(source_id, config)
    file_key = "fixture-file"
    evidence.append(
        _evidence(
            source_id,
            file_key,
            request_url=list_url,
            source_url=list_url,
            long_term_url=list_long_term,
            body=body,
            content_type="application/json",
            fetch_url=list_url,
            file_name=path.name,
            locator=path.name,
            rights_status=rights["rights_status"],
            attribution_required=rights["attribution_required"],
            attribution_text=rights["attribution_text"],
            rights_evidence_url=rights["rights_evidence_url"],
            metadata={**fixture_meta, "media_kind": "json"},
        )
    )
    for index, raw in enumerate(raw_items):
        item = dict(raw)
        item["source_id"] = source_id
        item["source_url"] = item.get("source_url") or list_url
        item["long_term_url"] = (
            item.get("long_term_url") or item["source_url"]
        )
        item["detail_url_status"] = item.get("detail_url_status") or "missing"
        item["rights_status"] = (
            item.get("rights_status") or rights["rights_status"]
        )
        item["attribution_required"] = item.get(
            "attribution_required", rights["attribution_required"]
        )
        item["attribution_text"] = (
            item.get("attribution_text") or rights["attribution_text"]
        )
        item["rights_evidence_url"] = (
            item.get("rights_evidence_url") or rights["rights_evidence_url"]
        )
        item["_evidence_key"] = file_key
        detail_html = item.get("detail_html") or item.get("body")
        if detail_html is not None:
            detail_key = f"item:{item.get('source_item_id', index)}"
            request_url = (
                item.get("fetched_url")
                or item.get("detail_url")
                or item["source_url"]
            )
            encoded = (
                detail_html.encode("utf-8")
                if isinstance(detail_html, str)
                else bytes(detail_html)
            )
            evidence.append(
                _evidence(
                    source_id,
                    detail_key,
                    request_url=request_url,
                    source_url=item["source_url"],
                    long_term_url=item["long_term_url"],
                    body=encoded,
                    content_type="text/html",
                    final_url=item.get("final_url"),
                    detail_url=item.get("detail_url"),
                    detail_url_status=item["detail_url_status"],
                    locator=item.get("locator") or f"{path.name}#items[{index}]",
                    file_name=f"{item.get('source_item_id', index)}.html",
                    rights_status=item["rights_status"],
                    attribution_required=item["attribution_required"],
                    attribution_text=item["attribution_text"],
                    rights_evidence_url=item["rights_evidence_url"],
                    metadata={**fixture_meta, "media_kind": "html"},
                )
            )
            item["_evidence_key"] = detail_key
        download = item.get("download")
        if download:
            download_body = download.get("body", "")
            encoded = (
                download_body.encode("utf-8")
                if isinstance(download_body, str)
                else bytes(download_body)
            )
            download_key = f"download:{item.get('source_item_id', index)}"
            download_name = download.get("file_name") or Path(
                urlparse(download.get("url") or "").path
            ).name
            evidence.append(
                _evidence(
                    source_id,
                    download_key,
                    request_url=download["url"],
                    source_url=download["url"],
                    long_term_url=item["long_term_url"],
                    body=encoded,
                    content_type=guess_content_type(download_name),
                    final_url=download.get("final_url"),
                    detail_url=item.get("detail_url"),
                    detail_url_status=item["detail_url_status"],
                    fetch_url=download["url"],
                    file_url=download.get("file_url") or download["url"],
                    locator=item.get("locator"),
                    file_name=download_name,
                    rights_status=item["rights_status"],
                    attribution_required=item["attribution_required"],
                    attribution_text=item["attribution_text"],
                    rights_evidence_url=item["rights_evidence_url"],
                    declared_size=download.get("declared_size"),
                    metadata={"media_kind": "attachment"},
                )
            )
        items.append(item)
    return DiscoveryResult(items=items, evidence=evidence)


def _discover_dialect_fixture(source_id, config):
    from .sources.dialect_org import parse_detail

    path = _fixture_path(config)
    if not path.exists():
        return DiscoveryResult(errors=[f"fixture not found: {path}"])
    body, fixture_meta = _fixture_metadata(path)
    payload = json.loads(body.decode("utf-8"))
    evidence = []
    items = []
    rights = _rights_fields(source_id, config)
    for index, raw in enumerate(payload.get("items") or []):
        item_id = str(raw["source_item_id"])
        parsed = parse_detail(item_id, raw["detail_html"], rights)
        point = parsed["point"]
        key = f"item:{item_id}"
        evidence.append(
            _evidence(
                source_id,
                key,
                request_url=point["source_url"],
                source_url=point["source_url"],
                long_term_url=point["long_term_url"],
                body=raw["detail_html"].encode("utf-8"),
                content_type="text/html",
                detail_url=point["detail_url"],
                detail_url_status=point["detail_url_status"],
                locator=f"{path.name}#items[{index}]",
                file_name=f"{item_id}.html",
                rights_status=point["rights_status"],
                attribution_required=point["attribution_required"],
                attribution_text=point["attribution_text"],
                rights_evidence_url=point["rights_evidence_url"],
                metadata={**fixture_meta, "media_kind": "html"},
            )
        )
        parsed["_evidence_key"] = key
        items.append(parsed)
    return DiscoveryResult(items=items, evidence=evidence)


def _discover_feiyiw_fixture(source_id, config):
    from .sources.feiyiw import parse_article_detail

    path = _fixture_path(config)
    if not path.exists():
        return DiscoveryResult(errors=[f"fixture not found: {path}"])
    body, fixture_meta = _fixture_metadata(path)
    payload = json.loads(body.decode("utf-8"))
    evidence = []
    items = []
    rights = _rights_fields(source_id, config)
    for index, raw in enumerate(payload.get("items") or []):
        item_id = str(raw["source_item_id"])
        fetched_url = raw.get("fetched_url") or (
            "https://www.feiyiw.cn/index.php?"
            f"app=article&act=view&article_id={item_id}"
        )
        record = parse_article_detail(raw["detail_html"], fetched_url, rights)
        key = f"item:{item_id}"
        evidence.append(
            _evidence(
                source_id,
                key,
                request_url=fetched_url,
                source_url=record["source_url"],
                long_term_url=record["long_term_url"],
                body=raw["detail_html"].encode("utf-8"),
                content_type="text/html",
                detail_url=record["detail_url"],
                detail_url_status=record["detail_url_status"],
                locator=f"{path.name}#items[{index}]",
                file_name=f"{item_id}.html",
                rights_status=record["rights_status"],
                attribution_required=record["attribution_required"],
                attribution_text=record["attribution_text"],
                rights_evidence_url=record["rights_evidence_url"],
                metadata={**fixture_meta, "media_kind": "html"},
            )
        )
        record["_evidence_key"] = key
        items.append(record)
    return DiscoveryResult(items=items, evidence=evidence)


def _discover_cass_fixture(source_id, config):
    path = _fixture_path(config)
    if not path.exists():
        return DiscoveryResult(errors=[f"fixture not found: {path}"])
    body, fixture_meta = _fixture_metadata(path)
    items = cass_linguistics.discover(
        config.fixture_root,
        getattr(config, "registry_path", None),
    )
    evidence = []
    for index, item in enumerate(items):
        item["source_id"] = source_id
        item["detail_url"] = item["source_url"]
        item["detail_url_status"] = "direct"
        item["publish_status"] = "draft"
        key = f"item:{item['source_item_id']}"
        item["_evidence_key"] = key
        evidence.append(
            _evidence(
                source_id,
                key,
                request_url=item["source_url"],
                source_url=item["source_url"],
                long_term_url=item["long_term_url"],
                body=_json_bytes(
                    {
                        "source_item_id": item["source_item_id"],
                        "title": item["title"],
                        "body_text": item.get("body_text"),
                    }
                ),
                content_type="application/json",
                detail_url=item["detail_url"],
                detail_url_status="direct",
                locator=f"{path.name}#items[{index}]",
                file_name=f"{item['source_item_id']}.json",
                rights_status=item["rights_status"],
                attribution_required=item["attribution_required"],
                attribution_text=item["attribution_text"],
                rights_evidence_url=item["rights_evidence_url"],
                metadata={**fixture_meta, "media_kind": "json"},
            )
        )
    return DiscoveryResult(items=items, evidence=evidence)


def _discover_saved_html(source_id, config):
    from bs4 import BeautifulSoup

    root = Path(config.source_root or "")
    patterns = {
        "nzrc": "job-detail-*.html",
        "nongmuhr": "employ-view-*.html",
    }
    paths = sorted(root.glob(patterns[source_id]))
    evidence = []
    items = []
    rights = _rights_fields(source_id, config)
    for path in paths:
        body = path.read_bytes()
        match = re.search(r"(\d+)", path.stem)
        source_item_id = match.group(1) if match else path.stem
        if source_id == "nzrc":
            source_url = (
                f"http://www.nzrc.com.cn/company/jobs/{source_item_id}.html"
            )
        else:
            source_url = (
                f"http://nongmuhr.com/EmployView.aspx?id={source_item_id}"
            )
        key = path.name
        evidence.append(
            _evidence(
                source_id,
                key,
                request_url=source_url,
                source_url=source_url,
                long_term_url=source_url,
                body=body,
                content_type="text/html",
                detail_url=source_url,
                detail_url_status="direct",
                locator=path.name,
                file_name=path.name,
                rights_status=rights["rights_status"],
                attribution_required=rights["attribution_required"],
                attribution_text=rights["attribution_text"],
                rights_evidence_url=rights["rights_evidence_url"],
                metadata={
                    "fixture_path": str(path),
                    "fixture_sha256": sha256_bytes(body),
                    "media_kind": "html",
                },
            )
        )
        try:
            soup = BeautifulSoup(
                body.decode("utf-8", errors="ignore"), "html.parser"
            )
            title_node = (
                soup.select_one(".job-name")
                or soup.select_one("#joblist")
                or soup.select_one("h1")
                or soup.title
            )
            title = (
                title_node.get_text(" ", strip=True)
                if title_node
                else path.stem
            )
            if source_id == "nongmuhr":
                title = title.split(" ")[0]
            text = soup.get_text(" ", strip=True)
            item = {
                "source_id": source_id,
                "source_item_id": source_item_id,
                "title": title,
                "company_name": None,
                "description": text[:2000],
                "source_url": source_url,
                "long_term_url": source_url,
                "detail_url": source_url,
                "detail_url_status": "direct",
                "publish_status": "draft",
                "locator": path.name,
                **rights,
                "extra": {
                    "rights_risk": {
                        "status": rights["rights_status"],
                        "notice": "只保存公开岗位摘要和来源链接。",
                    }
                },
                "_evidence_key": key,
            }
            if source_id == "nzrc":
                company = None
                for node in soup.select("a[href*='company']"):
                    value = node.get_text(" ", strip=True)
                    if value and len(value) < 80:
                        company = value
                        break
                item["company_name"] = company
            items.append(item)
        except Exception as exc:
            return DiscoveryResult(
                items=items,
                evidence=evidence,
                errors=[f"{path.name}: {exc}"],
            )
    return DiscoveryResult(items=items, evidence=evidence)


def _discover_aqsc_xlsx(source_id, config):
    from openpyxl import load_workbook

    root = Path(config.xlsx_root or "")
    evidence = []
    items = []
    rights = _rights_fields(source_id, config)
    for file_name, spec in _AQSC_FILES.items():
        path = root / file_name
        if not path.exists():
            return DiscoveryResult(
                items=items,
                evidence=evidence,
                errors=[f"missing xlsx fixture: {path}"],
            )
        body = path.read_bytes()
        key = file_name
        evidence.append(
            _evidence(
                source_id,
                key,
                request_url=spec["attachment_url"],
                source_url=spec["attachment_url"],
                long_term_url=spec["announcement_url"],
                body=body,
                content_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                fetch_url=spec["attachment_url"],
                locator=file_name,
                file_name=file_name,
                rights_status=rights["rights_status"],
                attribution_required=rights["attribution_required"],
                attribution_text=rights["attribution_text"],
                rights_evidence_url=rights["rights_evidence_url"],
                metadata={
                    "fixture_path": str(path),
                    "fixture_sha256": sha256_bytes(body),
                    "media_kind": "xlsx",
                },
            )
        )
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            header_index = None
            for index, candidate in enumerate(rows):
                names = {str(value or "").strip() for value in candidate}
                if "产品名称" in names or "产品全称" in names:
                    header_index = index
                    break
            if header_index is None:
                raise ValueError("missing product header")
            headers = [
                str(value or "").strip() for value in rows[header_index]
            ]
            for row_no, row in enumerate(
                rows[header_index + 1 :], start=header_index + 2
            ):
                values = dict(zip(headers, row))
                product = values.get("产品名称") or values.get("产品全称")
                if not product:
                    continue
                locator = f"{file_name}#{sheet.title}!A{row_no}"
                items.append(
                    {
                        "source_id": source_id,
                        "source_item_id": (
                            f"{spec['cert_type']}:{sheet.title}:{row_no}"
                        ),
                        "cert_type": spec["cert_type"],
                        "product_name": str(product).strip(),
                        "product_name_raw": str(product),
                        "normalized_name": str(product).strip().lower(),
                        "province": values.get("省份"),
                        "county": values.get("县域"),
                        "region": values.get("县域"),
                        "applicant": values.get("获证单位")
                        or values.get("申请单位全称"),
                        "certificate_no": values.get("证书编号"),
                        "source_url": spec["attachment_url"],
                        "long_term_url": spec["announcement_url"],
                        "detail_url": None,
                        "detail_url_status": "missing",
                        "fetch_url": spec["attachment_url"],
                        "publish_status": "blocked",
                        "locator": locator,
                        **rights,
                        "extra": {
                            "is_aquatic": None,
                            "announcement_url": spec["announcement_url"],
                            "attachment_url": spec["attachment_url"],
                        },
                        "_evidence_key": key,
                    }
                )
            workbook.close()
        except Exception as exc:
            return DiscoveryResult(
                items=items,
                evidence=evidence,
                errors=[f"{file_name}: {exc}"],
            )
    return DiscoveryResult(items=items, evidence=evidence)


def _cnipa_detail(path):
    for marker, url in _CNIPA_DETAIL_URLS.items():
        if marker in path.name:
            return marker, url
    return None


def _discover_cnipa_pdf(source_id, config):
    from pypdf import PdfReader

    root = Path(config.pdf_root or "")
    candidates = sorted(root.glob("*.pdf"))
    if not candidates:
        return DiscoveryResult(errors=[f"missing pdf fixture: {root}"])
    rights = _rights_fields(source_id, config)
    evidence = []
    items = []
    errors = []
    for path in candidates:
        mapping = _cnipa_detail(path)
        if mapping is None:
            errors.append(f"no documented detail URL mapping for {path.name}")
            continue
        announcement_no, detail_url = mapping
        body = path.read_bytes()
        key = f"pdf:{path.name}"
        evidence.append(
            _evidence(
                source_id,
                key,
                request_url=detail_url,
                source_url=detail_url,
                long_term_url=detail_url,
                body=body,
                content_type="application/pdf",
                fetch_url=detail_url,
                file_url=f"fixture://cnipa/{path.name}",
                locator=path.name,
                file_name=path.name,
                rights_status=rights["rights_status"],
                attribution_required=rights["attribution_required"],
                attribution_text=rights["attribution_text"],
                rights_evidence_url=rights["rights_evidence_url"],
                metadata={
                    "fixture_path": str(path),
                    "fixture_sha256": sha256_bytes(body),
                    "media_kind": "pdf",
                    "announcement_no": announcement_no,
                },
            )
        )
        try:
            reader = PdfReader(path)
            for page_no, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                if page_no < 2 or page_no > 28:
                    continue
                lines = [
                    line.strip()
                    for line in text.splitlines()
                    if line.strip()
                ]
                for index, line in enumerate(lines):
                    if line.isdigit() and index + 2 < len(lines):
                        product = lines[index + 1]
                        if product in {"产品名称", "序号"} or len(product) < 2:
                            continue
                        items.append(
                            {
                                "source_id": source_id,
                                "source_item_id": (
                                    f"cnipa:{announcement_no}:"
                                    f"{page_no}:{line}"
                                ),
                                "cert_type": "地理标志",
                                "product_name": product,
                                "product_name_raw": product,
                                "normalized_name": product.strip().lower(),
                                "province": (
                                    lines[index + 2]
                                    if index + 2 < len(lines)
                                    else None
                                ),
                                "region": None,
                                "applicant": (
                                    lines[index + 3]
                                    if index + 3 < len(lines)
                                    else None
                                ),
                                "origin_scope": " ".join(
                                    lines[index + 4 : index + 8]
                                ),
                                "source_url": detail_url,
                                "long_term_url": detail_url,
                                "detail_url": None,
                                "detail_url_status": "missing",
                                "fetch_url": detail_url,
                                "publish_status": "blocked",
                                "locator": (
                                    f"pdf:{path.name} "
                                    f"page={page_no} row={line}"
                                ),
                                **rights,
                                "extra": {
                                    "announcement_url": detail_url,
                                    "announcement_no": announcement_no,
                                    "pdf_page": page_no,
                                    "pdf_row": line,
                                },
                                "_evidence_key": key,
                            }
                        )
        except Exception as exc:
            errors.append(f"{path.name}: {exc}")
    return DiscoveryResult(items=items, evidence=evidence, errors=errors)


def _discover_cnipa_announcement_fixture(source_id, config):
    path = Path(config.fixture_root) / "announcements.json"
    if not path.exists():
        return DiscoveryResult(errors=[f"fixture not found: {path}"])
    body = path.read_bytes()
    payload = json.loads(body.decode("utf-8"))
    rows = payload.get("data", {}).get("data", [])
    rights = _rights_fields(source_id, config)
    key = path.name
    request_url = "https://ggfw.cnipa.gov.cn/dlbzsq/landmark/product/search/list"
    evidence = [
        _evidence(
            source_id,
            key,
            request_url=request_url,
            source_url=request_url,
            long_term_url="https://ggfw.cnipa.gov.cn/dlbzsq/dbQuery",
            body=body,
            content_type="application/json",
            locator=path.name,
            file_name=path.name,
            rights_status=rights["rights_status"],
            attribution_required=rights["attribution_required"],
            attribution_text=rights["attribution_text"],
            rights_evidence_url=rights["rights_evidence_url"],
            metadata={
                "fixture_path": str(path),
                "fixture_sha256": sha256_bytes(body),
                "media_kind": "json",
            },
        )
    ]
    items = []
    for row in rows:
        source_item_id = str(row["id"])
        detail_url = (
            f"https://ggfw.cnipa.gov.cn/dlbzsq/detail?id={source_item_id}"
        )
        items.append(
            {
                "source_id": source_id,
                "source_item_id": source_item_id,
                "title": row.get("name") or f"公告 {source_item_id}",
                "record_kind": "announcement",
                "source_url": request_url,
                "long_term_url": detail_url,
                "detail_url": detail_url,
                "detail_url_status": "direct",
                "publish_status": "blocked",
                "locator": f"announcement:{source_item_id}",
                **rights,
                "_evidence_key": key,
            }
        )
    return DiscoveryResult(items=items, evidence=evidence)


def _ihchina_list_url(page_no, limit):
    return (
        "https://www.ihchina.cn/getProject.html?"
        f"category_id=16&limit={limit}&p={page_no}"
    )


def _ihchina_items(page, page_no, limit, rights, fixture_meta, path):
    request_url = _ihchina_list_url(page_no, limit)
    items = []
    for raw in page.get("list") or []:
        source_item_id = str(raw["id"])
        detail_url = (
            f"https://www.ihchina.cn/project_details/{source_item_id}.html"
        )
        items.append(
            {
                "source_id": "ihchina",
                "source_item_id": source_item_id,
                "title": raw.get("title") or "",
                "ich_code": raw.get("num"),
                "project_serial": str(raw.get("auto_id") or ""),
                "source_project_num": raw.get("project_num"),
                "category": raw.get("type"),
                "project_type": raw.get("cate"),
                "announcement_batch": (raw.get("rx_time") or "").replace(
                    "</br>", " "
                ),
                "declaration_region": raw.get("unit") or raw.get("province"),
                "protection_unit": raw.get("protect_unit"),
                "description_html": raw.get("content"),
                "description_text": _strip_html(raw.get("content") or ""),
                "source_url": request_url,
                "long_term_url": detail_url,
                "detail_url": detail_url,
                "detail_url_status": "direct",
                "publish_status": "draft",
                "locator": f"list-page:{page_no}:id:{source_item_id}",
                **rights,
                "_evidence_key": f"page:{page_no}",
                "_fixture": {
                    **fixture_meta,
                    "fixture_path": str(path) if path else None,
                    "page": page_no,
                },
            }
        )
    return items


def _merge_ihchina_detail(item, detail):
    for key in (
        "title",
        "project_serial",
        "ich_code",
        "announcement_batch",
        "category",
        "province",
        "project_type",
        "declaration_region",
        "protection_unit",
        "description_html",
        "description_text",
    ):
        if detail.get(key):
            item[key] = detail[key]
    extra = dict(item.get("extra") or {})
    extra["detail"] = {
        "inheritors": detail.get("inheritors") or [],
        "related_projects": detail.get("related_projects") or [],
        "related_articles": detail.get("related_articles") or [],
        "media_links": detail.get("media_links") or [],
    }
    item["extra"] = extra


def _attach_ihchina_fixture_details(
    source_id,
    config,
    payload,
    items,
    evidence,
    rights,
    fixture_meta,
    path,
):
    from .sources import ihchina

    details = payload.get("details") or {}
    for item in items:
        item_id = item["source_item_id"]
        detail = details.get(item_id) or {}
        html = detail.get("html") or detail.get("detail_html")
        if html is None:
            html = item.get("detail_html")
        if html is None:
            continue
        detail_url = item["detail_url"]
        detail_key = f"detail:{item_id}"
        evidence.append(
            _evidence(
                source_id,
                detail_key,
                request_url=detail_url,
                final_url=detail.get("final_url"),
                source_url=detail_url,
                long_term_url=detail_url,
                body=html.encode("utf-8"),
                content_type="text/html",
                detail_url=detail_url,
                detail_url_status="direct",
                locator=f"project-detail:{item_id}",
                file_name=f"ihchina-project-{item_id}.html",
                rights_status=rights["rights_status"],
                attribution_required=rights["attribution_required"],
                attribution_text=rights["attribution_text"],
                rights_evidence_url=rights["rights_evidence_url"],
                metadata={**fixture_meta, "media_kind": "html"},
            )
        )
        parsed = ihchina.parse_project_detail(html, detail_url)
        inheritor_details = detail.get("inheritors") or {}
        for inheritor in parsed["inheritors"]:
            inheritor_id = inheritor.get("source_item_id")
            inheritor_html = inheritor_details.get(inheritor_id)
            if not inheritor_html:
                continue
            inheritor_url = inheritor["source_url"]
            evidence.append(
                _evidence(
                    source_id,
                    f"inheritor:{inheritor_id}",
                    request_url=inheritor_url,
                    source_url=inheritor_url,
                    long_term_url=inheritor_url,
                    body=inheritor_html.encode("utf-8"),
                    content_type="text/html",
                    detail_url=inheritor_url,
                    detail_url_status="direct",
                    locator=f"inheritor-detail:{inheritor_id}",
                    file_name=f"ihchina-inheritor-{inheritor_id}.html",
                    rights_status=rights["rights_status"],
                    attribution_required=rights["attribution_required"],
                    attribution_text=rights["attribution_text"],
                    rights_evidence_url=rights["rights_evidence_url"],
                    metadata={**fixture_meta, "media_kind": "html"},
                )
            )
            inheritor.update(
                ihchina.parse_inheritor_detail(
                    inheritor_html,
                    inheritor_url,
                )
            )
            inheritor["source_item_id"] = inheritor_id
        _merge_ihchina_detail(item, parsed)
        item["_evidence_key"] = detail_key


def _fetch_ihchina_detail(url):
    request = Request(
        url,
        headers={
            "User-Agent": (
                "yuexiang-crawler/0.1 (competition; noncommercial)"
            )
        },
    )
    with urlopen(request, timeout=30) as response:
        return (
            response.read(),
            getattr(response, "status", 200),
            response.geturl(),
            dict(response.headers.items()),
        )


def _attach_ihchina_live_details(source_id, items, evidence, rights):
    from .sources import ihchina

    errors = []
    for item in items:
        item_id = item["source_item_id"]
        detail_url = item["detail_url"]
        detail_key = f"detail:{item_id}"
        try:
            body, status_code, final_url, headers = _fetch_ihchina_detail(
                detail_url
            )
        except HTTPError as exc:
            body = exc.read()
            final_url = exc.geturl() or detail_url
            headers = dict(exc.headers.items())
            evidence.append(
                _evidence(
                    source_id,
                    detail_key,
                    request_url=detail_url,
                    final_url=final_url,
                    source_url=detail_url,
                    long_term_url=detail_url,
                    body=body,
                    content_type=exc.headers.get_content_type(),
                    detail_url=detail_url,
                    detail_url_status="direct",
                    locator=f"project-detail:{item_id}",
                    file_name=f"ihchina-project-{item_id}.html",
                    rights_status=rights["rights_status"],
                    attribution_required=rights["attribution_required"],
                    attribution_text=rights["attribution_text"],
                    rights_evidence_url=rights["rights_evidence_url"],
                    status_code=exc.code,
                    response_headers=headers,
                    fetch_status="http_error",
                    error_message=str(exc),
                    metadata={"media_kind": "html"},
                )
            )
            errors.append(f"project {item_id}: HTTP {exc.code}")
            continue

        evidence.append(
            _evidence(
                source_id,
                detail_key,
                request_url=detail_url,
                final_url=final_url,
                source_url=detail_url,
                long_term_url=detail_url,
                body=body,
                content_type=headers.get("Content-Type", "text/html"),
                detail_url=detail_url,
                detail_url_status="direct",
                locator=f"project-detail:{item_id}",
                file_name=f"ihchina-project-{item_id}.html",
                rights_status=rights["rights_status"],
                attribution_required=rights["attribution_required"],
                attribution_text=rights["attribution_text"],
                rights_evidence_url=rights["rights_evidence_url"],
                status_code=status_code,
                response_headers=headers,
                metadata={"media_kind": "html"},
            )
        )
        parsed = ihchina.parse_project_detail(
            body.decode("utf-8", errors="replace"),
            detail_url,
        )
        for inheritor in parsed["inheritors"]:
            inheritor_id = inheritor.get("source_item_id")
            inheritor_url = inheritor.get("source_url")
            if not inheritor_id or not inheritor_url:
                continue
            try:
                (
                    inheritor_body,
                    inheritor_status,
                    inheritor_final_url,
                    inheritor_headers,
                ) = _fetch_ihchina_detail(inheritor_url)
            except HTTPError as exc:
                inheritor_body = exc.read()
                inheritor_final_url = exc.geturl() or inheritor_url
                inheritor_headers = dict(exc.headers.items())
                evidence.append(
                    _evidence(
                        source_id,
                        f"inheritor:{inheritor_id}",
                        request_url=inheritor_url,
                        final_url=inheritor_final_url,
                        source_url=inheritor_url,
                        long_term_url=inheritor_url,
                        body=inheritor_body,
                        content_type=exc.headers.get_content_type(),
                        detail_url=inheritor_url,
                        detail_url_status="direct",
                        locator=f"inheritor-detail:{inheritor_id}",
                        file_name=f"ihchina-inheritor-{inheritor_id}.html",
                        rights_status=rights["rights_status"],
                        attribution_required=rights["attribution_required"],
                        attribution_text=rights["attribution_text"],
                        rights_evidence_url=rights[
                            "rights_evidence_url"
                        ],
                        status_code=exc.code,
                        response_headers=inheritor_headers,
                        fetch_status="http_error",
                        error_message=str(exc),
                        metadata={"media_kind": "html"},
                    )
                )
                errors.append(
                    f"inheritor {inheritor_id}: HTTP {exc.code}"
                )
                continue
            evidence.append(
                _evidence(
                    source_id,
                    f"inheritor:{inheritor_id}",
                    request_url=inheritor_url,
                    final_url=inheritor_final_url,
                    source_url=inheritor_url,
                    long_term_url=inheritor_url,
                    body=inheritor_body,
                    content_type=inheritor_headers.get(
                        "Content-Type", "text/html"
                    ),
                    detail_url=inheritor_url,
                    detail_url_status="direct",
                    locator=f"inheritor-detail:{inheritor_id}",
                    file_name=f"ihchina-inheritor-{inheritor_id}.html",
                    rights_status=rights["rights_status"],
                    attribution_required=rights["attribution_required"],
                    attribution_text=rights["attribution_text"],
                    rights_evidence_url=rights["rights_evidence_url"],
                    status_code=inheritor_status,
                    response_headers=inheritor_headers,
                    metadata={"media_kind": "html"},
                )
            )
            inheritor.update(
                ihchina.parse_inheritor_detail(
                    inheritor_body.decode("utf-8", errors="replace"),
                    inheritor_url,
                )
            )
            inheritor["source_item_id"] = inheritor_id
        _merge_ihchina_detail(item, parsed)
        item["_evidence_key"] = detail_key
    return errors


def _discover_ihchina_fixture(source_id, config):
    path = _fixture_path(config)
    if not path.exists():
        return DiscoveryResult(errors=[f"fixture not found: {path}"])
    body, fixture_meta = _fixture_metadata(path)
    payload = json.loads(body.decode("utf-8"))
    pages = payload.get("pages") or [payload]
    rights = _rights_fields(source_id, config)
    evidence = []
    items = []
    seen = set()
    for fallback_page, page in enumerate(pages, start=1):
        page_no = int(page.get("p") or fallback_page)
        request_url = _ihchina_list_url(page_no, config.page_limit)
        key = f"page:{page_no}"
        evidence.append(
            _evidence(
                source_id,
                key,
                request_url=request_url,
                source_url=request_url,
                long_term_url=request_url,
                body=_json_bytes(page),
                content_type="application/json",
                locator=f"page={page_no}",
                file_name=f"ihchina-page-{page_no}.json",
                rights_status=rights["rights_status"],
                attribution_required=rights["attribution_required"],
                attribution_text=rights["attribution_text"],
                rights_evidence_url=rights["rights_evidence_url"],
                metadata={
                    **fixture_meta,
                    "page": page_no,
                    "media_kind": "json",
                },
            )
        )
        for item in _ihchina_items(
            page, page_no, config.page_limit, rights, fixture_meta, path
        ):
            if item["source_item_id"] in seen:
                continue
            seen.add(item["source_item_id"])
            items.append(item)
    _attach_ihchina_fixture_details(
        source_id,
        config,
        payload,
        items,
        evidence,
        rights,
        fixture_meta,
        path,
    )
    return DiscoveryResult(items=items, evidence=evidence)


def _discover_ihchina_live(source_id, config):
    limit = int(config.page_limit)
    evidence = []
    items = []
    seen = set()
    rights = _rights_fields(source_id, config)
    page_no = 1
    total = None
    while True:
        request_url = _ihchina_list_url(page_no, limit)
        request = Request(
            request_url,
            headers={
                "User-Agent": (
                    "yuexiang-crawler/0.1 (competition; noncommercial)"
                )
            },
        )
        try:
            with urlopen(request, timeout=30) as response:
                body = response.read()
                status_code = getattr(response, "status", 200)
                final_url = response.geturl()
                headers = dict(response.headers.items())
            payload = json.loads(body.decode("utf-8"))
        except HTTPError as exc:
            body = exc.read()
            request_url = exc.geturl() or request_url
            evidence.append(
                _evidence(
                    source_id,
                    f"page:{page_no}",
                    request_url=request_url,
                    source_url=request_url,
                    long_term_url=request_url,
                    body=body,
                    content_type=exc.headers.get_content_type(),
                    locator=f"page={page_no}",
                    file_name=f"ihchina-page-{page_no}.html",
                    rights_status=rights["rights_status"],
                    attribution_required=rights["attribution_required"],
                    attribution_text=rights["attribution_text"],
                    rights_evidence_url=rights["rights_evidence_url"],
                    status_code=exc.code,
                    response_headers=dict(exc.headers.items()),
                    fetch_status="http_error",
                    error_message=str(exc),
                    metadata={"page": page_no, "media_kind": "html"},
                )
            )
            return DiscoveryResult(
                items=items,
                evidence=evidence,
                errors=[f"page {page_no}: HTTP {exc.code}"],
            )
        key = f"page:{page_no}"
        evidence.append(
            _evidence(
                source_id,
                key,
                request_url=request_url,
                source_url=request_url,
                long_term_url=final_url,
                body=body,
                content_type=headers.get("Content-Type", "application/json"),
                locator=f"page={page_no}",
                file_name=f"ihchina-page-{page_no}.json",
                rights_status=rights["rights_status"],
                attribution_required=rights["attribution_required"],
                attribution_text=rights["attribution_text"],
                rights_evidence_url=rights["rights_evidence_url"],
                status_code=status_code,
                response_headers=headers,
                metadata={"page": page_no, "media_kind": "json"},
            )
        )
        total = payload.get("total", total)
        for item in _ihchina_items(
            payload, page_no, limit, rights, {}, None
        ):
            if item["source_item_id"] in seen:
                continue
            seen.add(item["source_item_id"])
            item["source_url"] = request_url
            item["_evidence_key"] = key
            items.append(item)
        if not payload.get("list") or page_no * limit >= int(total or 0):
            break
        page_no += 1
    detail_errors = _attach_ihchina_live_details(
        source_id,
        items,
        evidence,
        rights,
    )
    return DiscoveryResult(
        items=items,
        evidence=evidence,
        errors=detail_errors,
    )


def _fiveajob_page_url(page_no):
    if page_no == 1:
        return "http://www.5ajob.com/job/"
    return f"http://www.5ajob.com/job/index-p{page_no}.html"


def _discover_fiveajob(source_id, config):
    from .sources import fiveajob

    path = _fixture_path(config)
    if not path.exists():
        return DiscoveryResult(errors=[f"fixture not found: {path}"])
    body, fixture_meta = _fixture_metadata(path)
    payload = json.loads(body.decode("utf-8"))
    pages = payload.get("pages") or [{"page": 1, "items": []}]
    rights = _rights_fields(source_id, config)
    evidence = []
    for fallback_page, page in enumerate(pages, start=1):
        page_no = int(page.get("page") or fallback_page)
        request_url = _fiveajob_page_url(page_no)
        evidence.append(
            _evidence(
                source_id,
                f"page:{page_no}",
                request_url=request_url,
                source_url=request_url,
                long_term_url=request_url,
                body=_json_bytes(page),
                content_type="application/json",
                locator=f"page={page_no}",
                file_name=f"5ajob-page-{page_no}.json",
                rights_status=rights["rights_status"],
                attribution_required=rights["attribution_required"],
                attribution_text=rights["attribution_text"],
                rights_evidence_url=rights["rights_evidence_url"],
                metadata={
                    **fixture_meta,
                    "page": page_no,
                    "media_kind": "json",
                },
            )
        )
    discovered = fiveajob.discover(config)
    items = discovered.get("items") or []
    for item in items:
        page_no = int(item.get("extra", {}).get("list_page") or 1)
        item["_evidence_key"] = f"page:{page_no}"
    return DiscoveryResult(items=items, evidence=evidence)


def _strip_html(value):
    from html import unescape

    return re.sub(r"<[^>]+>", "", unescape(value)).strip()


def _discover(source_id, mode, config):
    if source_id == "fixture-source":
        return _discover_generic_fixture(source_id, config)
    if source_id == "ihchina":
        if mode == "live":
            return _discover_ihchina_live(source_id, config)
        return _discover_ihchina_fixture(source_id, config)
    if source_id == "aqsc-moa":
        return _discover_aqsc_xlsx(source_id, config)
    if source_id == "cnipa":
        if mode == "announcement-fixture":
            return _discover_cnipa_announcement_fixture(source_id, config)
        return _discover_cnipa_pdf(source_id, config)
    if source_id in {"nzrc", "nongmuhr"}:
        return _discover_saved_html(source_id, config)
    if source_id == "dialect-org":
        return _discover_dialect_fixture(source_id, config)
    if source_id == "feiyiw":
        return _discover_feiyiw_fixture(source_id, config)
    if source_id == "cass-linguistics":
        return _discover_cass_fixture(source_id, config)
    if source_id == "5ajob":
        return _discover_fiveajob(source_id, config)
    raise ValueError(f"unsupported source: {source_id}")


def _stage_item(conn, source_id, mode, raw_fetch_id, item, retrieved_at):
    if source_id == "dialect-org":
        point = item["point"]
        lexeme = item["lexeme"]
        source_url = point["source_url"]
        long_term_url = point["long_term_url"]
        detail_url = point.get("detail_url")
        detail_url_status = point["detail_url_status"]
        fetch_url = point.get("fetch_url")
        locator = point.get("locator")
        rights_status = point["rights_status"]
        attribution_required = point.get("attribution_required")
        attribution_text = point.get("attribution_text")
        rights_evidence_url = point.get("rights_evidence_url")
        source_item_id = point["source_item_id"]
        title = lexeme["lemma"]
        extra = {
            "point": point.get("extra") or {},
            "lexeme": lexeme.get("extra") or {},
        }
    else:
        source_url = item.get("source_url")
        long_term_url = item.get("long_term_url")
        detail_url = item.get("detail_url")
        detail_url_status = item.get("detail_url_status")
        fetch_url = item.get("fetch_url")
        locator = item.get("locator")
        rights_status = item.get("rights_status")
        attribution_required = item.get("attribution_required")
        attribution_text = item.get("attribution_text")
        rights_evidence_url = item.get("rights_evidence_url")
        source_item_id = str(item.get("source_item_id") or raw_fetch_id)
        title = (
            item.get("title")
            or item.get("product_name")
            or item.get("name")
            or source_item_id
        )
        extra = item.get("extra") or {}
    validate_detail_locator(detail_url_status, detail_url, locator)
    content_hash = database.sha256_text(
        json.dumps(item, ensure_ascii=False, sort_keys=True, default=str)
    )
    conn.execute(
        """
        INSERT INTO stage_source_items (
            raw_fetch_id, source_id, source_item_id, source_url,
            long_term_url, detail_url, detail_url_status, fetch_url, locator,
            rights_status, attribution_required, attribution_text,
            rights_evidence_url, title, retrieved_at, content_hash,
            parser_version, extra
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(raw_fetch_id, source_item_id) DO UPDATE SET
            source_url=excluded.source_url,
            long_term_url=excluded.long_term_url,
            detail_url=excluded.detail_url,
            detail_url_status=excluded.detail_url_status,
            fetch_url=excluded.fetch_url,
            locator=excluded.locator,
            rights_status=excluded.rights_status,
            attribution_required=excluded.attribution_required,
            attribution_text=excluded.attribution_text,
            rights_evidence_url=excluded.rights_evidence_url,
            title=excluded.title,
            retrieved_at=excluded.retrieved_at,
            content_hash=excluded.content_hash,
            parser_version=excluded.parser_version,
            extra=excluded.extra
        """,
        (
            raw_fetch_id,
            source_id,
            source_item_id,
            source_url,
            long_term_url,
            detail_url,
            detail_url_status,
            fetch_url,
            locator,
            rights_status,
            1 if attribution_required else 0,
            attribution_text,
            rights_evidence_url,
            title,
            retrieved_at,
            content_hash,
            f"{source_id}-{mode}-v1",
            json.dumps(extra, ensure_ascii=False),
        ),
    )
    return content_hash


def _write_std(conn, source_id, item, retrieved_at, content_hash):
    if source_id == "dialect-org":
        point = dict(item["point"])
        lexeme = dict(item["lexeme"])
        point["content_hash"] = database.sha256_text(
            json.dumps(point, ensure_ascii=False, sort_keys=True, default=str)
        )
        lexeme["content_hash"] = content_hash
        database.upsert_std_dialect(
            conn, {"point": point, "lexeme": lexeme}, retrieved_at
        )
        return
    normal = dict(item)
    normal["content_hash"] = content_hash
    if source_id == "feiyiw":
        database.upsert_std_ich_article(conn, normal, retrieved_at)
    elif source_id == "cass-linguistics":
        database.upsert_std_research_article(conn, normal, retrieved_at)
    elif source_id == "ihchina":
        database.upsert_std_ich_project(conn, normal, retrieved_at)
    elif source_id in {"aqsc-moa", "cnipa"}:
        database.upsert_std_agri_product(conn, normal, retrieved_at)
    else:
        database.upsert_std_job(conn, normal, retrieved_at)


def run_source(source_id, mode, config):
    registry = source_registry.load_registry(
        getattr(config, "registry_path", None)
    )
    source = registry.get(source_id)
    if source is None and source_id not in _TEST_SOURCE_MODES:
        raise ValueError(f"unsupported source: {source_id}")
    run_id = (
        f"source:{source_id}:{mode}:"
        f"{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')}"
    )
    if source is not None:
        source_decision = source_registry.evaluate_source(source)
        if not source_decision.allowed:
            return SourceRunResult(
                run_id=run_id,
                source_id=source_id,
                status="blocked",
                discovered_count=0,
                fetched_count=0,
                written_count=0,
                skipped_count=0,
                failed_count=0,
                errors=(source_decision.reason or "source_blocked",),
            )
        allowed_modes = source_registry.run_modes(source)
    else:
        allowed_modes = _TEST_SOURCE_MODES[source_id]
    if mode not in allowed_modes:
        raise ValueError(f"mode {mode} is not implemented for {source_id}")

    retrieved_at = _now()
    try:
        discovery = _discover(source_id, mode, config)
    except Exception as exc:
        return SourceRunResult(
            run_id=run_id,
            source_id=source_id,
            status="failed",
            discovered_count=0,
            fetched_count=0,
            written_count=0,
            skipped_count=0,
            failed_count=1,
            errors=(str(exc),),
        )

    conn = database.connect(config.database_path)
    evidence_ids = {}
    try:
        for raw_evidence in discovery.evidence:
            evidence_ids[raw_evidence.key] = persist_raw_evidence(
                conn,
                raw_evidence,
                config.download_root,
                config.max_download_bytes,
                retrieved_at,
            )
        conn.commit()
    except Exception as exc:
        conn.rollback()
        conn.close()
        return SourceRunResult(
            run_id=run_id,
            source_id=source_id,
            status="failed",
            discovered_count=len(discovery.items),
            fetched_count=0,
            written_count=0,
            skipped_count=0,
            failed_count=max(1, len(discovery.items)),
            errors=tuple(discovery.errors + [str(exc)]),
        )

    written = 0
    skipped = 0
    failed = len(discovery.errors)
    errors = list(discovery.errors)
    try:
        for item in discovery.items:
            evidence_key = item.get("_evidence_key")
            raw_fetch_id = evidence_ids.get(evidence_key)
            if raw_fetch_id is None and evidence_ids:
                raw_fetch_id = next(iter(evidence_ids.values()))
            source_url = (
                item["point"].get("source_url")
                if source_id == "dialect-org"
                else item.get("source_url")
            )
            long_term_url = (
                item["point"].get("long_term_url")
                if source_id == "dialect-org"
                else item.get("long_term_url")
            )
            if raw_fetch_id is None or not source_url or not long_term_url:
                failed += 1
                errors.append(
                    f"{item.get('source_item_id')}: missing raw evidence or URL"
                )
                continue
            try:
                gate_row = (
                    item["point"] if source_id == "dialect-org" else item
                )
                try:
                    validate_detail_locator(
                        gate_row.get("detail_url_status"),
                        gate_row.get("detail_url"),
                        gate_row.get("locator"),
                    )
                except ValueError:
                    skipped += 1
                    continue
                content_hash = _stage_item(
                    conn,
                    source_id,
                    mode,
                    raw_fetch_id,
                    item,
                    retrieved_at,
                )
                if item.get("record_kind") == "announcement":
                    skipped += 1
                    continue
                decision = evaluate_publication(gate_row)
                if not decision.allowed:
                    skipped += 1
                    continue
                _write_std(conn, source_id, item, retrieved_at, content_hash)
                written += 1
            except Exception as exc:
                failed += 1
                errors.append(
                    f"{item.get('source_item_id', '?')}: {exc}"
                )
        conn.commit()
    except Exception as exc:
        conn.rollback()
        errors.append(str(exc))
        failed = max(failed, 1)
    finally:
        conn.close()

    discovered = len(discovery.items)
    if discovered and written + skipped + failed == discovered:
        if failed == 0:
            status = "succeeded"
        elif failed == discovered:
            status = "failed"
        else:
            status = "partial"
    elif discovered == 0:
        status = "failed" if errors else "partial"
    else:
        status = "partial"
    return SourceRunResult(
        run_id=run_id,
        source_id=source_id,
        status=status,
        discovered_count=discovered,
        fetched_count=len(evidence_ids),
        written_count=written,
        skipped_count=skipped,
        failed_count=failed,
        errors=tuple(errors),
    )
