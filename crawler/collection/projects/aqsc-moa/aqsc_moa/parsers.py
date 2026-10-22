"""Parsers for aqsc-moa XLSX attachments, embedded tables, and fixtures."""

from __future__ import annotations

import hashlib
import io
import json
import mimetypes
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from parsel import Selector


XLSX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

_HEADER_ALIASES = {
    "seq": ("序号",),
    "province": ("省份",),
    "county": ("县域",),
    "certificate_no": ("证书编号",),
    "product_name": ("产品名称", "产品全称"),
    "applicant": ("获证单位", "申请单位全称"),
    "production_scale_raw": ("生产规模",),
    "main_operators_raw": ("主要生产经营单位",),
}


def _text(value):
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def _normalized_header(value):
    return "".join(str(value or "").split())


def _header_positions(row):
    positions = {}
    source_names = {}
    for index, value in enumerate(row):
        header = _normalized_header(value)
        if not header:
            continue
        for field, aliases in _HEADER_ALIASES.items():
            matched = next(
                (
                    alias
                    for alias in aliases
                    if header == alias or header.startswith(alias)
                ),
                None,
            )
            if matched is not None and field not in positions:
                positions[field] = index
                source_names[field] = _text(value)
                break
    return positions, source_names


def _cell(row, positions, field):
    index = positions.get(field)
    if index is None or index >= len(row):
        return None
    return _text(row[index])


def _row_hash(payload):
    body = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(body).hexdigest()


def _key_part(value):
    return " ".join(str(value or "").split())


def _natural_key(
    *,
    certificate_no,
    batch,
    product_name,
    province,
    county,
    applicant,
):
    """Build a row identity that survives file and row reordering."""

    certificate_no = _key_part(certificate_no).upper().replace(" ", "")
    if certificate_no:
        return f"01:certificate={certificate_no}"
    return "02:" + "|".join(
        (
            f"batch={_key_part(batch)}",
            f"product={_key_part(product_name)}",
            f"province={_key_part(province)}",
            f"county={_key_part(county)}",
            f"applicant={_key_part(applicant)}",
        )
    )


def _announcement_id(metadata, fallback):
    return str(
        metadata.get("announcement_id")
        or metadata.get("announcement_no")
        or fallback
    )


def _record(
    *,
    metadata,
    announcement_id,
    product_name,
    province,
    county,
    certificate_no,
    applicant,
    production_scale_raw,
    main_operators_raw,
    locator,
    row_no,
    seq,
    source_names,
    source_kind,
    file_name=None,
    sheet=None,
    table_no=None,
    page_url=None,
):
    cert_type = (
        metadata.get("cert_type")
        or _infer_cert_type(file_name or "", metadata)
    )
    batch = (
        metadata.get("batch")
        or metadata.get("announcement_no")
        or announcement_id
    )
    extra = {
        "announcement_id": announcement_id,
        "announcement_no": metadata.get("announcement_no"),
        "batch": batch,
        "cert_type": cert_type,
        "product_name": product_name,
        "product_name_raw": product_name,
        "normalized_name": product_name,
        "province": province,
        "county": county,
        "certificate_no": certificate_no,
        "applicant": applicant,
        "production_scale_raw": production_scale_raw,
        "main_operators_raw": main_operators_raw,
        "is_aquatic": None,
        "file_name": file_name,
        "sheet": sheet,
        "table_no": table_no,
        "row_no": row_no,
        "seq": seq,
        "source_kind": source_kind,
        "source_field_names": source_names,
    }
    return {
        "source_item_id": _natural_key(
            certificate_no=certificate_no,
            batch=batch,
            product_name=product_name,
            province=province,
            county=county,
            applicant=applicant,
        ),
        "title": product_name,
        "source_url": (
            metadata.get("source_url") or page_url
        ),
        "long_term_url": (
            metadata.get("long_term_url") or page_url
        ),
        "detail_url": None,
        "detail_url_status": "missing",
        "locator": locator,
        "published_at": metadata.get("published_at"),
        "content_type": (
            XLSX_MIME_TYPE if source_kind == "xlsx" else "text/html"
        ),
        "content_hash": _row_hash(extra),
        "parser_version": (
            "aqsc-moa-xlsx-v2"
            if source_kind == "xlsx"
            else "aqsc-moa-html-v1"
        ),
        "candidate_status": "deferred",
        "candidate_reason": (
            metadata.get("candidate_reason") or "file-source:fixture"
        ),
        "extra": extra,
    }


def parse_xlsx_rows(body, *, file_name, metadata=None):
    """Parse product rows from a real or fixture XLSX attachment."""

    metadata = dict(metadata or {})
    workbook = load_workbook(
        io.BytesIO(bytes(body)), read_only=True, data_only=True
    )
    records = []
    seen = set()
    try:
        for sheet in workbook.worksheets:
            header_row = None
            positions = {}
            source_names = {}
            for row_no, row in enumerate(
                sheet.iter_rows(values_only=True), start=1
            ):
                candidate, candidate_names = _header_positions(row)
                if "seq" in candidate and "province" in candidate:
                    header_row = row_no
                    positions = candidate
                    source_names = candidate_names
                    break
            if header_row is None:
                continue

            for row_no, row in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                seq = _cell(row, positions, "seq")
                product_name = _cell(row, positions, "product_name")
                if not seq or not product_name:
                    continue
                announcement_id = _announcement_id(
                    metadata, Path(file_name).stem
                )
                locator = (
                    f"sheet={sheet.title}; row={row_no}; seq={seq}"
                )
                record = _record(
                    metadata=metadata,
                    announcement_id=announcement_id,
                    product_name=product_name,
                    province=_cell(row, positions, "province"),
                    county=_cell(row, positions, "county"),
                    certificate_no=_cell(
                        row, positions, "certificate_no"
                    ),
                    applicant=_cell(row, positions, "applicant"),
                    production_scale_raw=_cell(
                        row, positions, "production_scale_raw"
                    ),
                    main_operators_raw=_cell(
                        row, positions, "main_operators_raw"
                    ),
                    locator=locator,
                    row_no=row_no,
                    seq=seq,
                    source_names=source_names,
                    source_kind="xlsx",
                    file_name=file_name,
                    sheet=sheet.title,
                )
                if record["source_item_id"] in seen:
                    continue
                seen.add(record["source_item_id"])
                records.append(record)
    finally:
        workbook.close()
    return records


def parse_embedded_rows(html, *, page_url, metadata=None):
    """Parse product rows from aqsc-moa `.Details_con table` HTML."""

    metadata = dict(metadata or {})
    selector = Selector(text=html)
    fallback_id = Path(urlparse(page_url).path).stem or page_url
    announcement_id = _announcement_id(metadata, fallback_id)
    records = []
    seen = set()
    tables = selector.css(".Details_con table")
    for table_no, table in enumerate(tables, start=1):
        rows = []
        for row in table.css("tr"):
            values = [
                _text(cell.xpath("string(.)").get())
                for cell in row.css("th, td")
            ]
            rows.append(values)

        header_row = None
        positions = {}
        source_names = {}
        for row_no, row in enumerate(rows, start=1):
            candidate, candidate_names = _header_positions(row)
            if "seq" in candidate and "province" in candidate:
                header_row = row_no
                positions = candidate
                source_names = candidate_names
                break
        if header_row is None:
            continue

        for row_no, row in enumerate(
            rows[header_row:], start=header_row + 1
        ):
            seq = _cell(row, positions, "seq")
            product_name = _cell(row, positions, "product_name")
            if not seq or not product_name:
                continue
            locator = (
                f"table={table_no}; row={row_no}; seq={seq}"
            )
            record = _record(
                metadata=metadata,
                announcement_id=announcement_id,
                product_name=product_name,
                province=_cell(row, positions, "province"),
                county=_cell(row, positions, "county"),
                certificate_no=_cell(
                    row, positions, "certificate_no"
                ),
                applicant=_cell(row, positions, "applicant"),
                production_scale_raw=_cell(
                    row, positions, "production_scale_raw"
                ),
                main_operators_raw=_cell(
                    row, positions, "main_operators_raw"
                ),
                locator=locator,
                row_no=row_no,
                seq=seq,
                source_names=source_names,
                source_kind="embedded_html",
                table_no=table_no,
                page_url=page_url,
            )
            if record["source_item_id"] in seen:
                continue
            seen.add(record["source_item_id"])
            records.append(record)
    return records


def _infer_cert_type(file_name, metadata):
    text = " ".join(
        [
            str(file_name or ""),
            str(metadata.get("title") or ""),
            str(metadata.get("long_term_url") or ""),
        ]
    )
    if "特质" in text or "tznp" in text.lower():
        return "tezhi_nongpin"
    return "mingteyouxin"


def fixture_files(root):
    records = []
    for path in sorted(Path(root).glob("*.xlsx")):
        body = path.read_bytes()
        records.append(
            {
                "file_name": path.name,
                "mime_type": mimetypes.guess_type(path.name)[0]
                or XLSX_MIME_TYPE,
                "byte_size": len(body),
                "sha256": hashlib.sha256(body).hexdigest(),
                "locator": "sheet=Sheet1; row=3",
            }
        )
    return records


def fixture_item(path):
    path = Path(path)
    source_url = (
        "http://www.aqsc.agri.cn/tzgg/202511/t20251128_468194.htm"
    )
    return {
        "source_id": "aqsc-moa",
        "source_item_id": path.stem,
        "title": path.stem,
        "source_url": source_url,
        "long_term_url": source_url,
        "detail_url": None,
        "detail_url_status": "missing",
        "locator": "sheet=Sheet1; row=3",
        "run_id": "fixture",
        "mode": "fixture",
        "candidate_status": "deferred",
        "candidate_reason": "file-source:fixture",
    }
